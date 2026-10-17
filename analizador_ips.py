import argparse
import ipaddress
import json
import os
import sys
import time
from datetime import datetime, timedelta
from pathlib import Path

import pandas as pd
import requests

try:
    import truststore
    truststore.inject_into_ssl()
except ImportError:
    pass

ABUSEIPDB_URL = "https://api.abuseipdb.com/api/v2/check"
CACHE_FILE = Path("cache_abuseipdb.json")
HISTORIAL_FILE = Path("historial_ips.json")
WHITELIST_FILE = Path("whitelist.txt")
CONFIG_FILE = Path("config.json")

MIN_REPORTES_MALICIOSA = 2
MIN_CONFIDENCE_MALICIOSA = 25
DIAS_HISTORIAL = 90

COLUMNAS_IP = ["srcip", "data.srcip", "data.src_ip", "source.ip", "src_ip"]
COLUMNAS_PAIS = ["srccountry", "data.srccountry", "GeoLocation.country_name",
                 "geoip.country_name", "source.geo.country_name"]
COLUMNAS_REGLA = ["rule.description", "rule_description", "description"]
COLUMNAS_URL = ["url", "data.url", "data.uri", "http.url", "request_url"]
COLUMNAS_COUNT = ["count", "Count", "doc_count", "Cantidad", "Recuento"]
COLUMNAS_AGENTE = ["agent.name", "agent_name", "agent.hostname", "hostname"]

PAISES = {
    "US": "United States", "RU": "Russia", "CN": "China", "DE": "Germany",
    "FR": "France", "GB": "United Kingdom", "NL": "Netherlands", "BR": "Brazil",
    "IN": "India", "SG": "Singapore", "HK": "Hong Kong", "JP": "Japan",
    "KR": "South Korea", "VN": "Vietnam", "ID": "Indonesia", "TH": "Thailand",
    "TR": "Turkey", "IR": "Iran", "UA": "Ukraine", "PL": "Poland",
    "RO": "Romania", "BG": "Bulgaria", "LT": "Lithuania", "MD": "Moldova",
    "SC": "Seychelles", "IE": "Ireland", "ES": "Spain", "IT": "Italy",
    "PT": "Portugal", "BE": "Belgium", "SE": "Sweden", "NO": "Norway",
    "FI": "Finland", "CH": "Switzerland", "AT": "Austria", "CZ": "Czechia",
    "CA": "Canada", "MX": "Mexico", "AR": "Argentina", "CL": "Chile",
    "CO": "Colombia", "PE": "Peru", "VE": "Venezuela", "EC": "Ecuador",
    "UY": "Uruguay", "PY": "Paraguay", "BO": "Bolivia", "PA": "Panama",
    "AU": "Australia", "NZ": "New Zealand", "ZA": "South Africa",
    "EG": "Egypt", "NG": "Nigeria", "KE": "Kenya", "MA": "Morocco",
    "IL": "Israel", "SA": "Saudi Arabia", "AE": "United Arab Emirates",
    "PK": "Pakistan", "BD": "Bangladesh", "TW": "Taiwan", "MY": "Malaysia",
    "PH": "Philippines", "KZ": "Kazakhstan", "BY": "Belarus", "RS": "Serbia",
    "HR": "Croatia", "SI": "Slovenia", "SK": "Slovakia", "HU": "Hungary",
    "GR": "Greece", "DK": "Denmark", "EE": "Estonia", "LV": "Latvia",
    "CR": "Costa Rica", "DO": "Dominican Republic", "GT": "Guatemala",
    "HN": "Honduras", "SV": "El Salvador", "NI": "Nicaragua", "CU": "Cuba",
}


def cargar_api_key():
    key = os.environ.get("ABUSEIPDB_API_KEY", "").strip()
    if not key and CONFIG_FILE.exists():
        try:
            key = json.loads(CONFIG_FILE.read_text(encoding="utf-8")).get("abuseipdb_api_key", "").strip()
        except (json.JSONDecodeError, UnicodeDecodeError):
            pass
    if not key:
        sys.exit("ERROR: falta la API key. Defini ABUSEIPDB_API_KEY o crea config.json")
    return key


def cargar_whitelist():
    wl = {"exactas": set(), "prefijos": [], "redes": []}
    if not WHITELIST_FILE.exists():
        print(f"AVISO: no se encontro {WHITELIST_FILE}, se continua sin whitelist.")
        return wl
    for linea in WHITELIST_FILE.read_text(encoding="utf-8").splitlines():
        linea = linea.split("#")[0].strip()
        if not linea:
            continue
        if "/" in linea:
            try:
                wl["redes"].append(ipaddress.ip_network(linea, strict=False))
                continue
            except ValueError:
                print(f"AVISO: rango CIDR invalido en whitelist ignorado: {linea!r}")
                continue
        if linea.endswith("."):
            wl["prefijos"].append(linea)
            continue
        try:
            ipaddress.ip_address(linea)
            wl["exactas"].add(linea)
        except ValueError:
            print(f"AVISO: linea invalida en whitelist ignorada: {linea!r}")
    return wl


def esta_en_whitelist(ip, wl):
    if ip in wl["exactas"]:
        return True
    if any(ip.startswith(p) for p in wl["prefijos"]):
        return True
    if wl["redes"]:
        ip_obj = ipaddress.ip_address(ip)
        if any(ip_obj in red for red in wl["redes"]):
            return True
    return False


def detectar_columna(df, candidatas):
    def normalizar(nombre):
        n = nombre.lower().strip()
        for sufijo in (": descending", ": ascending", " descending", " ascending"):
            if n.endswith(sufijo):
                n = n[: -len(sufijo)].strip().rstrip(":").strip()
        return n

    columnas_norm = {normalizar(c): c for c in df.columns}
    for cand in candidatas:
        if cand.lower() in columnas_norm:
            return columnas_norm[cand.lower()]
    return None


def cargar_cache(dias_validez):
    if not CACHE_FILE.exists():
        return {}
    try:
        cache = json.loads(CACHE_FILE.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, UnicodeDecodeError):
        return {}
    limite = datetime.now() - timedelta(days=dias_validez)
    return {
        ip: datos for ip, datos in cache.items()
        if datetime.fromisoformat(datos.get("consultado", "1970-01-01")) > limite
    }


def guardar_cache(cache):
    CACHE_FILE.write_text(json.dumps(cache, indent=1, ensure_ascii=False), encoding="utf-8")


def cargar_historial():
    if not HISTORIAL_FILE.exists():
        return {}
    try:
        historial = json.loads(HISTORIAL_FILE.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, UnicodeDecodeError):
        print(f"AVISO: {HISTORIAL_FILE} esta danado, se empieza un historial nuevo.")
        return {}
    limite = (datetime.now() - timedelta(days=DIAS_HISTORIAL)).strftime("%Y-%m-%d")
    depurado = {}
    for ip, vistas in historial.items():
        vigentes = [v for v in vistas if v.get("fecha", "") >= limite]
        if vigentes:
            depurado[ip] = vigentes
    return depurado


def guardar_historial(historial):
    HISTORIAL_FILE.write_text(json.dumps(historial, indent=1, ensure_ascii=False), encoding="utf-8")


def describir_antecedentes(ip, historial, cliente, hoy):
    previas = [v for v in historial.get(ip, [])
               if not (v.get("fecha") == hoy and v.get("cliente") == cliente)]
    if not previas:
        return "Primera vez"
    previas.sort(key=lambda v: v.get("fecha", ""))
    ultima = previas[-1]
    fecha = datetime.strptime(ultima["fecha"], "%Y-%m-%d").strftime("%d/%m/%Y")
    clientes = sorted({v.get("cliente", "") for v in previas})
    veces = len(previas)
    texto_veces = "1 vez" if veces == 1 else f"{veces} veces"
    return (f"Ya registrada: {texto_veces} - ultima {fecha} "
            f"({ultima.get('resultado', '-')}) - clientes: {', '.join(clientes)}")


def registrar_en_historial(historial, ip, cliente, hoy, resultado):
    vistas = [v for v in historial.get(ip, [])
              if not (v.get("fecha") == hoy and v.get("cliente") == cliente)]
    vistas.append({"fecha": hoy, "cliente": cliente, "resultado": resultado})
    historial[ip] = vistas


def consultar_abuseipdb(ip, api_key):
    resp = requests.get(
        ABUSEIPDB_URL,
        headers={"Key": api_key, "Accept": "application/json"},
        params={"ipAddress": ip, "maxAgeInDays": 90},
        timeout=15,
    )
    if resp.status_code == 429:
        raise RuntimeError("Cuota diaria de AbuseIPDB agotada (HTTP 429).")
    resp.raise_for_status()
    d = resp.json()["data"]
    return {
        "reportes": d.get("totalReports", 0),
        "confidence": d.get("abuseConfidenceScore", 0),
        "isp": d.get("isp") or "Desconocido",
        "dominio": d.get("domain") or "",
        "uso": d.get("usageType") or "",
        "pais_abuse": d.get("countryCode") or "",
        "ultimo": d.get("lastReportedAt") or "",
        "consultado": datetime.now().isoformat(timespec="seconds"),
    }


def es_maliciosa(info):
    return (info["reportes"] >= MIN_REPORTES_MALICIOSA
            or info["confidence"] >= MIN_CONFIDENCE_MALICIOSA)


def veredicto(info):
    reportes = f"{info['reportes']:,}".replace(",", ".")
    detalle = f"({info['confidence']}% - {reportes} reportes)"
    if es_maliciosa(info):
        return f"Cuenta con registros maliciosos {detalle}"
    return f"No cuenta con registros maliciosos {detalle}"


def describir_ultimo_reporte(info):
    if not info or not info.get("ultimo"):
        return "-"
    try:
        fecha = datetime.fromisoformat(info["ultimo"].replace("Z", "+00:00"))
    except ValueError:
        return "-"
    dias = (datetime.now(fecha.tzinfo) - fecha).days
    if dias < 0:
        dias = 0
    if dias <= 30:
        return f"Reciente (hace {dias} dias)" if dias != 1 else "Reciente (hace 1 dia)"
    meses = dias // 30
    if meses < 12:
        return f"Antiguo (hace {meses} meses)" if meses != 1 else "Antiguo (hace 1 mes)"
    anios = dias // 365
    return f"Antiguo (hace {anios} anios)" if anios != 1 else "Antiguo (hace 1 anio)"


def main():
    parser = argparse.ArgumentParser(description="Analiza IPs de exports de Wazuh contra AbuseIPDB.")
    parser.add_argument("--csv", nargs="+", default=[])
    parser.add_argument("--txt", nargs="+", default=[])
    parser.add_argument("--cliente", required=True)
    parser.add_argument("--dias-cache", type=int, default=2)
    parser.add_argument("--salida", default=None)
    parser.add_argument("--detalle", action="store_true")
    args = parser.parse_args()

    if not args.csv and not args.txt:
        sys.exit("ERROR: indica al menos un archivo con --csv (exports de Wazuh) o --txt (lista de IPs)")

    api_key = cargar_api_key()
    whitelist = cargar_whitelist()
    print(f"Whitelist cargada: {len(whitelist['exactas'])} IPs exactas, "
          f"{len(whitelist['prefijos'])} prefijos, {len(whitelist['redes'])} rangos CIDR")

    registros = {}
    total_eventos = 0

    for archivo in args.csv:
        try:
            df = pd.read_csv(archivo)
        except Exception as e:
            sys.exit(f"ERROR leyendo {archivo}: {e}")

        col_ip = detectar_columna(df, COLUMNAS_IP)
        if not col_ip:
            print(f"AVISO: {archivo} no tiene columna de IP reconocible "
                  f"(columnas: {list(df.columns)}). Se omite este archivo.")
            continue
        col_pais = detectar_columna(df, COLUMNAS_PAIS)
        col_regla = detectar_columna(df, COLUMNAS_REGLA)
        col_url = detectar_columna(df, COLUMNAS_URL)
        col_count = detectar_columna(df, COLUMNAS_COUNT)
        col_agente = detectar_columna(df, COLUMNAS_AGENTE)
        print(f"{archivo}: {len(df)} filas | IP: {col_ip} | Pais: {col_pais} "
              f"| Regla: {col_regla} | URL: {col_url} | Count: {col_count} | Agente: {col_agente}")

        df[col_ip] = df[col_ip].astype(str).str.strip()

        for _, fila in df.iterrows():
            ip = fila[col_ip]
            try:
                ipaddress.ip_address(ip)
            except ValueError:
                continue
            reg = registros.setdefault(ip, {
                "eventos": 0,
                "pais": "",
                "reglas": set(),
                "urls": set(),
                "origenes": set(),
                "agentes": set(),
            })
            eventos_fila = 1
            if col_count and pd.notna(fila.get(col_count)):
                try:
                    eventos_fila = max(int(float(fila[col_count])), 1)
                except (ValueError, TypeError):
                    pass
            reg["eventos"] += eventos_fila
            total_eventos += eventos_fila
            reg["origenes"].add(Path(archivo).stem)
            if col_pais and pd.notna(fila.get(col_pais)) and not reg["pais"]:
                reg["pais"] = str(fila[col_pais]).strip()
            if col_regla and pd.notna(fila.get(col_regla)):
                reg["reglas"].add(str(fila[col_regla]).strip())
            if col_url and pd.notna(fila.get(col_url)):
                reg["urls"].add(str(fila[col_url]).strip())
            if col_agente and pd.notna(fila.get(col_agente)):
                reg["agentes"].add(str(fila[col_agente]).strip())

    for archivo in args.txt:
        ruta = Path(archivo)
        if not ruta.exists():
            sys.exit(f"ERROR: no se encontro {archivo}")
        validas = 0
        for linea in ruta.read_text(encoding="utf-8").splitlines():
            linea = linea.split("#")[0].strip().rstrip(",;")
            if not linea:
                continue
            try:
                ipaddress.ip_address(linea)
            except ValueError:
                print(f"AVISO: linea invalida en {archivo} ignorada: {linea!r}")
                continue
            reg = registros.setdefault(linea, {
                "eventos": 0,
                "pais": "",
                "reglas": set(),
                "urls": set(),
                "origenes": set(),
                "agentes": set(),
            })
            reg["eventos"] += 1
            total_eventos += 1
            reg["origenes"].add(ruta.stem)
            validas += 1
        print(f"{archivo}: {validas} IPs leidas")

    ips_unicas = list(registros.keys())
    print(f"Eventos: {total_eventos} | IPs unicas: {len(ips_unicas)}")

    cache = cargar_cache(args.dias_cache)
    historial = cargar_historial()
    hoy = datetime.now().strftime("%Y-%m-%d")
    consultas_nuevas = 0
    con_antecedentes = []
    filas_salida = []

    for i, ip in enumerate(sorted(ips_unicas), 1):
        reg = registros[ip]
        en_whitelist = esta_en_whitelist(ip, whitelist)
        es_privada = not ipaddress.ip_address(ip).is_global

        if en_whitelist or es_privada:
            info = None
        elif ip in cache:
            info = cache[ip]
        else:
            print(f"[{i}/{len(ips_unicas)}] consultando {ip} ...")
            try:
                info = consultar_abuseipdb(ip, api_key)
            except RuntimeError as e:
                print(f"AVISO: {e} Se marcan las restantes como 'Sin analizar'.")
                info = None
            except requests.RequestException as e:
                print(f"AVISO: error consultando {ip}: {e}")
                info = None
            else:
                cache[ip] = info
                consultas_nuevas += 1
                time.sleep(0.3)

        if en_whitelist or es_privada:
            antecedentes = "-"
        else:
            antecedentes = describir_antecedentes(ip, historial, args.cliente, hoy)
            if antecedentes != "Primera vez":
                con_antecedentes.append((ip, antecedentes))
            resultado = ("maliciosa" if es_maliciosa(info) else "limpia") if info else "sin analizar"
            registrar_en_historial(historial, ip, args.cliente, hoy, resultado)

        pais_fallback = ""
        if info and info.get("pais_abuse"):
            cod = info["pais_abuse"].strip().upper()
            pais_fallback = PAISES.get(cod, cod)

        filas_salida.append({
            "IP": ip,
            "Pais": (reg["pais"] or pais_fallback) or "Desconocido",
            "Eventos": reg["eventos"],
            "Servidor atacado": "; ".join(sorted(reg["agentes"]))[:120],
            "Reglas detectadas": "; ".join(sorted(reg["reglas"]))[:250],
            "URLs atacadas": "; ".join(sorted(reg["urls"]))[:250],
            "Tabla origen": "; ".join(sorted(reg["origenes"])),
            "Whitelist": "SI - EXCLUIDA" if en_whitelist else "No",
            "Reportes (AbuseIPDB)": info["reportes"] if info else "-",
            "Confidence": info["confidence"] if info else "-",
            "Analisis": ("Whitelist - no aplica" if en_whitelist
                         else "IP privada/interna - revisar origen" if es_privada
                         else (veredicto(info) if info else "Sin analizar")),
            "Servicio (ISP)": info["isp"] if info else "-",
            "Ultimo reporte": describir_ultimo_reporte(info) if not (en_whitelist or es_privada) else "-",
            "Antecedentes": antecedentes,
            "Tipo de uso": info["uso"] if info else "-",
        })

    guardar_cache(cache)
    guardar_historial(historial)
    print(f"Consultas nuevas a AbuseIPDB: {consultas_nuevas} (resto desde cache/whitelist)")

    if con_antecedentes:
        print(f"\nIPs ya registradas en analisis anteriores: {len(con_antecedentes)}")
        for ip, texto in con_antecedentes:
            print(f"  - {ip}: {texto}")

    df_out = pd.DataFrame(filas_salida)
    df_out["_orden"] = df_out["Analisis"].str.startswith("Cuenta con registros maliciosos")
    df_out = df_out.sort_values(["_orden", "Eventos"], ascending=[False, False]).drop(columns="_orden")

    columnas = ["IP", "Pais", "Analisis", "Servicio (ISP)", "Ultimo reporte", "Antecedentes"]
    if args.detalle:
        columnas += ["Eventos", "Servidor atacado", "Reglas detectadas", "URLs atacadas",
                     "Whitelist", "Reportes (AbuseIPDB)", "Confidence", "Tipo de uso",
                     "Tabla origen"]
    df_reporte = df_out[columnas]

    maliciosas = df_out[df_out["Analisis"].str.startswith("Cuenta con registros maliciosos")]
    df_bloqueo = maliciosas[["IP"]].rename(columns={"IP": "IPs a bloquear"}).reset_index(drop=True)

    fecha = datetime.now().strftime("%Y%m%d")
    salida = Path(args.salida) if args.salida else Path(
        f"reporte_{args.cliente.replace(' ', '_').lower()}_{fecha}.xlsx")

    with pd.ExcelWriter(salida, engine="openpyxl") as writer:
        df_reporte.to_excel(writer, sheet_name="Analisis IPs", index=False)
        df_bloqueo.to_excel(writer, sheet_name="IPs bloqueo", index=False)

        from openpyxl.styles import Font, PatternFill
        for nombre_hoja in writer.sheets:
            ws = writer.sheets[nombre_hoja]
            for celda in ws[1]:
                celda.font = Font(bold=True, color="FFFFFF")
                celda.fill = PatternFill("solid", fgColor="1F4E79")
            for col in ws.columns:
                ancho = max(len(str(c.value)) for c in col if c.value is not None)
                ws.column_dimensions[col[0].column_letter].width = min(ancho + 3, 60)

        ws = writer.sheets["Analisis IPs"]
        col_analisis = [c.value for c in ws[1]].index("Analisis") + 1
        rojo = PatternFill("solid", fgColor="FCE4E4")
        for fila in ws.iter_rows(min_row=2):
            if str(fila[col_analisis - 1].value or "").startswith("Cuenta con registros maliciosos"):
                for celda in fila:
                    celda.fill = rojo

    print(f"\nListo -> {salida}")
    print(f"IPs maliciosas detectadas: {len(maliciosas)}")
    if len(maliciosas):
        print("Recomendadas para bloqueo:")
        for ip in maliciosas["IP"]:
            print(f"  - {ip}")


if __name__ == "__main__":
    main()
